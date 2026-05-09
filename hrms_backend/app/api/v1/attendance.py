from datetime import date, datetime
from typing import Optional, Sequence

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ForbiddenError, NotFoundError
from app.core.permissions import require_permission
from app.dependencies import get_current_user, get_db
from app.models import Attendance, AttendanceRegularization, LeaveRequest, User
from app.schemas.attendance import (
    AttendanceCalendar,
    AttendanceCheckIn,
    AttendanceCheckOut,
    AttendanceOut,
    AttendanceSummary,
    AttendanceUpdate,
    AttendanceUploadResult,
    AttendanceUpsertByDate,
    RegularizationCreate,
    RegularizationDecision,
    RegularizationOut,
)
from app.services import attendance_service as att_svc
from app.services import employee_service as emp_svc
from app.services.storage import upload_fileobj

router = APIRouter(prefix="/attendance", tags=["attendance"])


async def _self_employee(db: AsyncSession, user: User):
    emp = await emp_svc.get_self_employee(db, user)
    if not emp:
        raise NotFoundError("No employee record linked to this account")
    return emp


@router.post("/check-in", response_model=AttendanceOut)
async def check_in(
    payload: AttendanceCheckIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = await _self_employee(db, user)
    att, _geo = await att_svc.check_in(
        db,
        emp,
        lat=payload.lat,
        lng=payload.lng,
        work_mode=payload.work_mode,
        selfie_url=payload.selfie_url,
        notes=payload.notes,
    )
    return att


@router.post("/check-in/with-selfie", response_model=AttendanceOut)
async def check_in_with_selfie(
    lat: float = Form(...),
    lng: float = Form(...),
    work_mode: str = Form("office"),
    notes: Optional[str] = Form(None),
    selfie: Optional[UploadFile] = File(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Multipart variant: upload the selfie photo as part of the request."""
    emp = await _self_employee(db, user)

    selfie_url: Optional[str] = None
    if selfie is not None:
        contents = await selfie.read()
        if contents:
            import io
            _, selfie_url = upload_fileobj(
                io.BytesIO(contents),
                key_prefix=f"attendance/{emp.employee_code}",
                filename=selfie.filename or "selfie.jpg",
                content_type=selfie.content_type,
            )

    att, _geo = await att_svc.check_in(
        db,
        emp,
        lat=lat,
        lng=lng,
        work_mode=work_mode,
        selfie_url=selfie_url,
        notes=notes,
    )
    return att


@router.post("/check-out", response_model=AttendanceOut)
async def check_out(
    payload: AttendanceCheckOut,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = await _self_employee(db, user)
    return await att_svc.check_out(db, emp, lat=payload.lat, lng=payload.lng)


@router.get("/me", response_model=list[AttendanceOut])
async def my_attendance(
    from_date: Optional[date] = None,
    to_date: Optional[date] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = await _self_employee(db, user)
    return await att_svc.list_for_employee(
        db, emp.id, from_date=from_date, to_date=to_date
    )


@router.get("/me/summary", response_model=AttendanceSummary)
async def my_summary(
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    today = date.today()
    emp = await _self_employee(db, user)
    return await att_svc.monthly_summary(
        db, emp.id, year=year or today.year, month=month or today.month
    )


async def _synthesize_leave_rows(
    db: AsyncSession,
    *,
    from_date: date,
    to_date: date,
    employee_ids: Optional[list[int]] = None,
):
    """Produce virtual Attendance rows (status='leave') for each
    (employee, date) covered by an approved LeaveRequest in the range.
    Used to surface employees on leave even when they didn't punch in."""
    from datetime import timedelta as _td

    stmt = (
        select(LeaveRequest)
        .where(LeaveRequest.status == "approved")
        .where(LeaveRequest.from_date <= to_date)
        .where(LeaveRequest.to_date >= from_date)
    )
    if employee_ids is not None:
        if not employee_ids:
            return []
        stmt = stmt.where(LeaveRequest.employee_id.in_(employee_ids))
    leaves = (await db.execute(stmt)).scalars().all()

    out: list[Attendance] = []
    for lr in leaves:
        d = max(lr.from_date, from_date)
        end = min(lr.to_date, to_date)
        while d <= end:
            virt = Attendance(
                employee_id=lr.employee_id,
                date=d,
                status="leave",
                work_mode="office",
                source="leave",
                total_minutes=0,
                notes=(lr.reason or None),
            )
            # Use a deterministic negative id derived from the leave row
            # so React's `key={r.id}` stays unique across the synthesized
            # set. Real (positive) attendance ids are never overwritten.
            virt.id = -(lr.id * 100 + (d - lr.from_date).days)
            out.append(virt)
            d = d + _td(days=1)
    return out


def _merge_attendance_with_leave(
    real: list[Attendance], virtual: list[Attendance]
) -> list[Attendance]:
    """Add synthesized leave rows for any (employee_id, date) that
    isn't already represented in the real attendance set."""
    seen = {(a.employee_id, a.date) for a in real}
    merged = list(real)
    for v in virtual:
        if (v.employee_id, v.date) not in seen:
            merged.append(v)
    merged.sort(key=lambda r: (r.date, r.employee_id), reverse=True)
    return merged


@router.get("/team", response_model=list[AttendanceOut])
async def team_attendance(
    on_date: Optional[date] = Query(default=None, alias="date"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    allowed = await emp_svc.visible_employee_ids(db, user)
    if allowed is not None and not allowed:
        return []
    on = on_date or date.today()
    stmt = select(Attendance).where(Attendance.date == on)
    if allowed is not None:
        stmt = stmt.where(Attendance.employee_id.in_(allowed))
    real = (await db.execute(stmt)).scalars().all()

    virtual = await _synthesize_leave_rows(
        db,
        from_date=on,
        to_date=on,
        employee_ids=list(allowed) if allowed is not None else None,
    )
    return _merge_attendance_with_leave(list(real), virtual)


@router.get("/report", response_model=list[AttendanceOut])
async def attendance_report(
    from_date: date,
    to_date: date,
    employee_id: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("attendance.view.all")),
):
    stmt = (
        select(Attendance)
        .where(Attendance.date >= from_date)
        .where(Attendance.date <= to_date)
        .order_by(Attendance.date.desc())
    )
    if employee_id:
        stmt = stmt.where(Attendance.employee_id == employee_id)
    real = (await db.execute(stmt)).scalars().all()

    virtual = await _synthesize_leave_rows(
        db,
        from_date=from_date,
        to_date=to_date,
        employee_ids=[employee_id] if employee_id else None,
    )
    return _merge_attendance_with_leave(list(real), virtual)


# ── Regularization ──────────────────────────────────────────────────────
@router.post(
    "/regularize",
    response_model=RegularizationOut,
    status_code=status.HTTP_201_CREATED,
)
async def request_regularization(
    payload: RegularizationCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = await _self_employee(db, user)
    return await att_svc.request_regularization(
        db,
        emp,
        on_date=payload.date,
        requested_check_in=payload.requested_check_in,
        requested_check_out=payload.requested_check_out,
        reason=payload.reason,
    )


@router.get("/regularize/pending", response_model=list[RegularizationOut])
async def pending_regularizations(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("approval.act")),
):
    return await att_svc.list_pending_regularizations(db)


@router.post("/regularize/{reg_id}/approve", response_model=RegularizationOut)
async def approve_regularization(
    reg_id: int,
    payload: RegularizationDecision,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission("approval.act")),
):
    return await att_svc.approve_regularization(
        db, reg_id, approver=user, decision_comment=payload.decision_comment
    )


@router.post("/regularize/{reg_id}/reject", response_model=RegularizationOut)
async def reject_regularization(
    reg_id: int,
    payload: RegularizationDecision,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission("approval.act")),
):
    return await att_svc.reject_regularization(
        db, reg_id, approver=user, decision_comment=payload.decision_comment
    )


# ── Admin: edit a single record ─────────────────────────────────────
@router.put("/{attendance_id}", response_model=AttendanceOut)
async def update_attendance(
    attendance_id: int,
    payload: AttendanceUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("attendance.view.all")),
):
    from app.models import Attendance as AttModel

    att = (
        await db.execute(select(AttModel).where(AttModel.id == attendance_id))
    ).scalar_one_or_none()
    if not att:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Attendance record not found")
    data = payload.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(att, k, v)
    if att.check_in and att.check_out:
        att.total_minutes = int(
            (att.check_out - att.check_in).total_seconds() // 60
        )
    await db.commit()
    await db.refresh(att)
    return att


# ── Calendar (self) ────────────────────────────────────────────────
@router.get("/me/calendar", response_model=AttendanceCalendar)
async def my_calendar(
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    today = date.today()
    emp = await _self_employee(db, user)
    return await att_svc.calendar_for_employee(
        db,
        employee_id=emp.id,
        year=year or today.year,
        month=month or today.month,
    )


# ── Calendar (admin / manager scoped) ──────────────────────────────
@router.get(
    "/employee/{employee_id}/calendar",
    response_model=AttendanceCalendar,
)
async def employee_calendar(
    employee_id: int,
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("attendance.view.all")),
):
    today = date.today()
    return await att_svc.calendar_for_employee(
        db,
        employee_id=employee_id,
        year=year or today.year,
        month=month or today.month,
    )


# ── Admin upsert single day from calendar editor ───────────────────
@router.put(
    "/employee/{employee_id}/by-date",
    response_model=AttendanceOut,
)
async def upsert_attendance_by_date(
    employee_id: int,
    payload: AttendanceUpsertByDate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission("attendance.view.all")),
):
    return await att_svc.upsert_by_date(
        db,
        employee_id=employee_id,
        on_date=payload.date,
        status_=payload.status,
        check_in=payload.check_in,
        check_out=payload.check_out,
        work_mode=payload.work_mode,
        notes=payload.notes,
        actor_user_id=user.id,
    )


# ── Admin: bulk upload from CSV / Excel ─────────────────────────────
@router.post("/upload", response_model=AttendanceUploadResult)
async def upload_attendance(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("attendance.view.all")),
):
    """Bulk-import attendance rows. Accepts CSV (preferred) or XLSX.

    Required columns: `employee_code`, `date` (YYYY-MM-DD; / and -
    separators also work). Optional columns: `status` (present, absent,
    half_day, wfh, leave, holiday), `check_in`, `check_out` (HH:MM or
    full ISO timestamp), `work_mode` (office, wfh, hybrid, field),
    `notes`. Existing (employee, date) rows are updated in place.

    Returns a summary plus a structured list of per-row errors so the
    admin can fix the spreadsheet and re-upload.
    """
    contents = await file.read()
    if not contents:
        from app.core.exceptions import ValidationError
        raise ValidationError("Uploaded file is empty")

    fname = (file.filename or "").lower()
    rows: list[dict] = []

    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            from app.core.exceptions import ValidationError
            raise ValidationError(
                "Excel uploads need `openpyxl` installed; please convert "
                "to CSV and try again."
            )
        import io as _io
        wb = load_workbook(filename=_io.BytesIO(contents), data_only=True)
        ws = wb.active
        header = None
        for r in ws.iter_rows(values_only=True):
            cells = [(str(c).strip() if c is not None else "") for c in r]
            if header is None:
                header = [h.lower() for h in cells]
                continue
            row = {header[i]: cells[i] if i < len(cells) else "" for i in range(len(header))}
            rows.append(row)
    else:
        import csv as _csv
        import io as _io
        text = contents.decode("utf-8-sig", errors="replace")
        reader = _csv.DictReader(_io.StringIO(text))
        if reader.fieldnames is None:
            from app.core.exceptions import ValidationError
            raise ValidationError("CSV is missing a header row")
        # Lowercase header keys so the parser is forgiving.
        norm_headers = [h.strip().lower() for h in reader.fieldnames]
        for r in reader:
            rows.append(
                {norm_headers[i]: (list(r.values())[i] or "")
                 for i in range(len(norm_headers))}
            )

    return await att_svc.bulk_import(db, rows)


@router.get("/upload/template.csv")
async def upload_template(
    _: User = Depends(require_permission("attendance.view.all")),
):
    """Download a starter CSV with the headers Jaikalki expects."""
    from fastapi.responses import StreamingResponse

    csv = (
        "employee_code,date,status,check_in,check_out,work_mode,notes\n"
        "JKT001,2026-05-01,present,09:30,18:15,office,\n"
        "JKT001,2026-05-02,half_day,10:30,15:00,office,Late grace\n"
        "JKT002,2026-05-01,wfh,09:45,18:30,wfh,Working from home\n"
    )
    return StreamingResponse(
        iter([csv]),
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="attendance-template.csv"',
        },
    )


# ── Admin: report CSV export ────────────────────────────────────────
@router.get("/report.csv")
async def attendance_report_csv(
    from_date: date,
    to_date: date,
    employee_id: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_permission("attendance.view.all")),
):
    import csv as _csv
    import io as _io
    from fastapi.responses import StreamingResponse

    from app.models import Employee as EmpModel

    stmt = (
        select(Attendance)
        .where(Attendance.date >= from_date)
        .where(Attendance.date <= to_date)
        .order_by(Attendance.date.desc(), Attendance.employee_id)
    )
    if employee_id:
        stmt = stmt.where(Attendance.employee_id == employee_id)
    real = (await db.execute(stmt)).scalars().all()
    virtual = await _synthesize_leave_rows(
        db,
        from_date=from_date,
        to_date=to_date,
        employee_ids=[employee_id] if employee_id else None,
    )
    rows = _merge_attendance_with_leave(list(real), virtual)
    emps = (await db.execute(select(EmpModel))).scalars().all()
    by_id = {e.id: e for e in emps}

    buf = _io.StringIO()
    w = _csv.writer(buf)
    w.writerow([
        "date", "employee_code", "name", "status", "work_mode",
        "check_in", "check_out", "minutes", "notes",
    ])
    for r in rows:
        e = by_id.get(r.employee_id)
        w.writerow([
            r.date.isoformat(),
            e.employee_code if e else "",
            f"{e.first_name} {e.last_name}" if e else "",
            r.status,
            r.work_mode,
            r.check_in.isoformat() if r.check_in else "",
            r.check_out.isoformat() if r.check_out else "",
            r.total_minutes,
            (r.notes or "").replace("\n", " "),
        ])
    buf.seek(0)
    fname = f"attendance-{from_date}-{to_date}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
